"""Read the single 'MATCH Sales Dashboard' master workbook and return everything the race
board needs, so the whole dashboard builds from one uploaded file.

Sheets used:
  Sales and count   store/day sales + transaction counts (this year)   -> stores, daily, tx
  Staff Daily sales staff/day sales + counts                           -> users, staff metrics (g7, ATV)
  TY Sales Target   target, 31-day target curve, Prize (ceiling), Area Leader, Region, Name
  LY Sales          per-branch per-day sales last year                 -> lyG / company ly
  Footfall          per-store visitors (period total)                  -> conversion

read_master() returns (stores, daily, targets, asof, extra) where stores/daily/targets match
readers.read_pos + read_targets, and `extra` carries the Phase-2 enrichment race_data consumes.
When the live Odoo API replaces the sales sheets, only this reader changes; race_data does not.
"""
import re
import datetime
import openpyxl

from .config import FX

_DAY_TXT = re.compile(r"^(\d{2}) (\w{3}) (\d{4})$")             # "01 Oct 2026"
PRIZE_MAP = {"car": "A", "smart phone": "B", "i pad": "C", "ipad": "C",
             "smart watch": "D", "gold bar": "E", "none": None, "": None}


def _day(cell):
    """A day-of-month int from a header cell that is a datetime or a 'DD Mon YYYY' string, else None."""
    if isinstance(cell, datetime.datetime):
        return cell.day
    if isinstance(cell, str):
        m = _DAY_TXT.match(cell.strip())
        if m:
            return int(m.group(1))
    return None


def _num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def read_master(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    # ---- Sales and count: store/day sales + transaction counts -------------------------------
    ws = wb["Sales and count"]
    rows = list(ws.iter_rows(values_only=True))
    datehdr = rows[1]                                           # row 2 holds the day headers
    pairs = [(j, _day(datehdr[j])) for j in range(6, len(datehdr), 2) if _day(datehdr[j])]
    stores, daily, tx, country = {}, {}, {}, {}
    year = month = None
    for j, _ in pairs:                                          # learn year/month from any real datetime header
        if isinstance(datehdr[j], datetime.datetime):
            year, month = datehdr[j].year, datehdr[j].month
    if year is None:
        year, month = 2026, 10                                  # text headers only: fall back to the program month
    day_total = {}
    for r in rows[3:]:
        code = r[0]
        if not isinstance(code, str) or not code.strip():
            continue
        code = code.strip().upper()
        c = (str(r[1]).strip() if r[1] else "KSA")
        country[code] = c if c in FX else "KSA"          # FX keys are KSA/UAE/Oman/Kuwait/Bahrain/Qatar
        mtd = tx_mtd = 0.0
        for j, dom in pairs:
            v, c = _num(r[j]), _num(r[j + 1])
            mtd += v; tx_mtd += c
            iso = datetime.date(year, month, dom).isoformat()
            daily.setdefault(iso, {})[code] = round(daily.get(iso, {}).get(code, 0) + v, 2)
            day_total[dom] = day_total.get(dom, 0) + v
        stores[code] = {"sales": round(mtd, 2), "users": []}
        tx[code] = round(tx_mtd)
    lastday = max((d for d, v in day_total.items() if v > 0), default=max(d for _, d in pairs))
    asof = datetime.date(year, month, lastday).isoformat()
    daily = {d: v for d, v in daily.items() if int(d[-2:]) <= lastday}   # drop empty trailing days

    # ---- Staff Daily sales: per-staff mtd sales, counts, and daily sales (for g7) -------------
    ws = wb["Staff Daily sales"]
    srows = list(ws.iter_rows(values_only=True))
    spairs = [(j, _day(srows[1][j])) for j in range(6, len(srows[1]), 2) if _day(srows[1][j])]
    staff = {}
    for r in srows[2:]:
        raw, store = r[3], r[1]
        if not isinstance(raw, str) or not isinstance(store, str) or raw.strip() in ("", "Total"):
            continue
        raw, store = raw.strip(), store.strip().upper()
        sday, stot, scount = {}, 0.0, 0.0
        for j, dom in spairs:
            v, c = _num(r[j]), _num(r[j + 1])
            if dom <= lastday:
                sday[dom] = v; stot += v; scount += c
        staff[raw] = {"store": store, "sales": round(stot, 2), "tx": round(scount), "daily": sday}
        if store in stores:
            stores[store]["users"].append([raw, round(stot, 2)])

    # ---- TY Sales Target: target, curve, ceiling (Prize), area leader, region, name ----------
    ws = wb["TY Sales Target"]
    trows = list(ws.iter_rows(values_only=True))
    tcols = [j for j in range(10, len(trows[1])) if isinstance(trows[1][j], datetime.datetime)]
    targets, ceiling, managers, names = {}, {}, {}, {}
    comp_daily = [_num(trows[2][j]) for j in tcols]             # row 3 = company daily target totals
    for r in trows[4:]:
        code = r[1]
        if not isinstance(code, str) or not code.strip():
            continue
        code = code.strip().upper()
        tgt = sum(_num(r[j]) for j in tcols)                   # October target = sum of the daily cells
        targets[code] = {"country": country.get(code, "KSA"), "region": (r[4] or "").strip() or "—",
                         "city": (r[3] or "").strip(), "target": round(tgt, 2)}
        names[code] = (str(r[2]).strip() if r[2] else code)
        ceiling[code] = PRIZE_MAP.get(str(r[8] or "").strip().lower())
        if r[5]:
            managers[code] = str(r[5]).strip()
    comp_tot = sum(comp_daily) or 1.0
    curve, run = [], 0.0
    for v in comp_daily:                                        # cumulative % of month target, per day
        run += v
        curve.append(round(run / comp_tot * 100, 2))

    # ---- LY Sales: per-branch per-day sales last year ----------------------------------------
    ly = {}
    ws = wb["LY Sales "]
    lrows = list(ws.iter_rows(values_only=True))
    lcols = [(j, _day(lrows[1][j])) for j in range(2, len(lrows[1])) if _day(lrows[1][j])]
    for r in lrows[4:]:
        code = r[1]
        if not isinstance(code, str) or not code.strip():
            continue
        ly[code.strip().upper()] = {dom: _num(r[j]) for j, dom in lcols}

    # ---- Footfall: period-total visitors per store -------------------------------------------
    foot = {}
    ws = wb["Footfall"]
    frows = list(ws.iter_rows(values_only=True))
    for r in frows[4:]:
        code = r[1]
        if not isinstance(code, str) or not code.strip():
            continue
        foot[code.strip().upper()] = sum(_num(c) for c in r[2:])

    extra = {"names": names, "ceiling": ceiling, "managers": managers, "curve": curve,
             "tx": tx, "footfall": foot, "ly": ly, "staff": staff}
    return stores, daily, targets, asof, extra
